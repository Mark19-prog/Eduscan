from __future__ import annotations

import os
import sys
import tempfile
import unittest
import io
import zipfile
import hashlib
import uuid
from datetime import date, datetime, time, timedelta
from pathlib import Path
from unittest.mock import patch


BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))
TEST_DIR = tempfile.TemporaryDirectory(prefix="eduscan-smoke-")
database_path = Path(TEST_DIR.name) / "smoke.db"
os.environ["DATABASE_URL"] = f"sqlite:///{database_path.as_posix()}"
os.environ["EDUSCAN_DATA_DIR"] = TEST_DIR.name
os.environ["EDUSCAN_SECRET_KEY"] = "smoke-test-secret-not-for-deployment"
os.environ["SMS_GATEWAY_ENABLED"] = "false"
os.environ["ATTENDANCE_SCHEDULER_ENABLED"] = "false"
os.environ["SF2_TEMPLATE_PATH"] = str(BACKEND_DIR / "data" / "templates" / "School Form 2 (SF2) Daily Attendance Report of Learners.xlsx")

from fastapi.testclient import TestClient  # noqa: E402
from app.main import app  # noqa: E402


class EduScanSmokeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.client_context = TestClient(app)
        cls.client = cls.client_context.__enter__()
        response = cls.client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
        assert response.status_code == 200, response.text
        cls.headers = {"Authorization": f"Bearer {response.json()['access_token']}"}

    @classmethod
    def tearDownClass(cls) -> None:
        cls.client_context.__exit__(None, None, None)
        from app.database import engine
        engine.dispose()
        TEST_DIR.cleanup()

    def test_biometric_enrollment_crud_and_audit(self) -> None:
        person = self.client.post("/api/persons", headers=self.headers, json={
            "external_id": "BIO-CRUD-001", "lrn": "999999999991", "full_name": "AUDIT, FACE TEST",
            "sex": "Female", "role": "Student", "grade": "10", "section": "Rizal",
            "guardian_phone": "+639171111111", "biometric_consent": True,
        })
        self.assertEqual(person.status_code, 200, person.text)
        person_id = person.json()["id"]

        import cv2
        import numpy as np
        from app.services.biometrics import ProcessedFace, biometric_service

        def face_set(offset: int) -> list[ProcessedFace]:
            result = []
            for index in range(15):
                image = np.full((200, 200), 90 + offset + index, dtype=np.uint8)
                cv2.circle(image, (100, 95), 58, 150 + index, -1)
                cv2.circle(image, (77, 79), 7, 15 + offset, -1)
                cv2.circle(image, (123, 79), 7, 15 + offset, -1)
                cv2.line(image, (78, 126), (122, 126 + (index % 3)), 25, 4)
                ok, encoded = cv2.imencode(".png", image)
                self.assertTrue(ok)
                result.append(ProcessedFace(image=image, encoded=encoded.tobytes(), quality=92 + index / 10))
            return result

        uploads = [("frames", (f"frame-{index}.jpg", b"test-frame", "image/jpeg")) for index in range(15)]
        with patch.object(biometric_service, "process_frame", side_effect=face_set(0)):
            created = self.client.post(f"/api/biometrics/enrollments/{person_id}", headers=self.headers,
                                       data={"reason": "Initial authorized enrollment"}, files=uploads)
        self.assertEqual(created.status_code, 200, created.text)
        self.assertEqual(created.json()["action"], "created")
        self.assertEqual(created.json()["enrollment"]["sample_count"], 15)
        from app.database import SessionLocal
        from app.models import BiometricModel, BiometricSample
        with SessionLocal() as db:
            self.assertEqual(db.query(BiometricSample).filter_by(person_id=person_id).count(), 15)
            self.assertEqual(db.query(BiometricModel).count(), 1)
            first_model_path = Path(db.query(BiometricModel).one().file_path)
            self.assertTrue(first_model_path.exists())

        listing = self.client.get("/api/biometrics/enrollments", headers=self.headers)
        self.assertEqual(listing.status_code, 200, listing.text)
        self.assertEqual(listing.json()[0]["external_id"], "BIO-CRUD-001")
        detail = self.client.get(f"/api/biometrics/enrollments/{person_id}", headers=self.headers)
        self.assertEqual(detail.status_code, 200, detail.text)
        self.assertEqual(len(detail.json()["samples"]), 15)
        self.assertNotIn("encrypted_path", detail.json()["samples"][0])
        self.assertNotIn("sha256", detail.json()["samples"][0])

        with patch.object(biometric_service, "process_frame", side_effect=face_set(20)):
            updated = self.client.put(f"/api/biometrics/enrollments/{person_id}", headers=self.headers,
                                      data={"reason": "Appearance changed after verification"}, files=uploads)
        self.assertEqual(updated.status_code, 200, updated.text)
        self.assertEqual(updated.json()["action"], "updated")
        with SessionLocal() as db:
            self.assertEqual(db.query(BiometricSample).filter_by(person_id=person_id).count(), 15)
            self.assertEqual(db.query(BiometricModel).count(), 1)
            replacement_model_path = Path(db.query(BiometricModel).one().file_path)
            self.assertTrue(replacement_model_path.exists())
        self.assertFalse(first_model_path.exists())

        deleted = self.client.request("DELETE", f"/api/biometrics/enrollments/{person_id}", headers=self.headers,
                                      json={"reason": "Authorized retention period ended"})
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertEqual(deleted.json()["removed_samples"], 15)
        self.assertEqual(deleted.json()["removed_model_records"], 1)
        self.assertIsNone(deleted.json()["model"])
        self.assertEqual(self.client.get("/api/biometrics/enrollments", headers=self.headers).json(), [])
        with SessionLocal() as db:
            self.assertEqual(db.query(BiometricSample).filter_by(person_id=person_id).count(), 0)
            self.assertEqual(db.query(BiometricModel).count(), 0)
        self.assertFalse(replacement_model_path.exists())

        events = self.client.get(f"/api/biometrics/enrollments/audit?person_id={person_id}", headers=self.headers)
        self.assertEqual(events.status_code, 200, events.text)
        self.assertEqual([item["action"] for item in events.json()], ["deleted", "updated", "created"])
        self.assertEqual(events.json()[0]["before"]["sample_count"], 15)
        self.assertEqual(events.json()[0]["after"]["sample_count"], 0)

    def test_persistent_attendance_correction_and_gradebook(self) -> None:
        person = self.client.post("/api/persons", headers=self.headers, json={
            "external_id": "SMOKE-001", "lrn": "123456789012", "full_name": "TEST, LEARNER A",
            "sex": "Male", "role": "Student", "grade": "10", "section": "Rizal",
            "guardian_phone": "+639171234567", "biometric_consent": True,
        })
        self.assertEqual(person.status_code, 200, person.text)
        person_id = person.json()["id"]

        schedule = self.client.post("/api/schedules", headers=self.headers, json={
            "grade": "10", "section": "Rizal", "subject": "Mathematics", "teacher_name": "Teacher Account",
            "weekdays": "0,1,2,3,4", "start_time": "07:30", "end_time": "08:30",
            "late_grace_minutes": 15, "active": True,
        })
        self.assertEqual(schedule.status_code, 200, schedule.text)

        today = date(2026, 8, 3).isoformat()
        correction = self.client.post("/api/attendance/corrections", headers=self.headers, json={
            "person_id": person_id, "attendance_date": today, "status": "Late",
            "time_in": "08:02", "time_out": None, "reason": "Verified signed gate log",
        })
        self.assertEqual(correction.status_code, 200, correction.text)
        rows = self.client.get(f"/api/attendance?date={today}", headers=self.headers)
        self.assertEqual(rows.status_code, 200, rows.text)
        corrected_row = next(item for item in rows.json() if item["person_id"] == person_id)
        self.assertEqual(corrected_row["status"], "Late")

        sf2 = self.client.get("/api/sf2/export?year=2026&month=8&grade=10&section=Rizal&school_id=SMOKE&school_year=2026-2027", headers=self.headers)
        self.assertEqual(sf2.status_code, 200, sf2.text)
        with zipfile.ZipFile(io.BytesIO(sf2.content)) as archive:
            self.assertIn("xl/media/image2.wmf", archive.namelist())
            self.assertIn("xl/media/eduscan-tardy.png", archive.namelist())
            drawing = archive.read("xl/drawings/drawing1.xml")
            self.assertIn(b"EduScan Tardy 14-4", drawing)

        class_key = "10-rizal-mathematics"
        gradebook = self.client.put(f"/api/gradebook/{class_key}", headers=self.headers, json={
            "class_key": class_key, "passing_grade": 75,
            "components": [
                {"id": "quiz-1", "category": "Quizzes", "label": "Quiz 1", "weight": 30, "max_score": 100},
                {"id": "summative-1", "category": "Summative Tests", "label": "Summative 1", "weight": 50, "max_score": 100},
                {"id": "periodic-1", "category": "Periodic Tests", "label": "Periodic 1", "weight": 20, "max_score": 100},
            ],
            "scores": {str(person_id): {"quiz-1": 88, "summative-1": 90, "periodic-1": 86}},
        })
        self.assertEqual(gradebook.status_code, 200, gradebook.text)
        loaded = self.client.get(f"/api/gradebook/{class_key}", headers=self.headers)
        self.assertEqual(loaded.status_code, 200, loaded.text)
        self.assertEqual(len(loaded.json()["components"]), 3)
        self.assertEqual(len(loaded.json()["scores"][str(person_id)]), 3)

        # Exercise actual OpenCV LBPH training and encrypted model persistence.
        import cv2
        import numpy as np
        from app.database import SessionLocal
        from app.models import BiometricModel, BiometricSample
        from app.services.biometrics import biometric_service
        with SessionLocal() as db:
            for index in range(15):
                image = np.zeros((200, 200), dtype=np.uint8)
                cv2.circle(image, (100, 95), 64, 130 + index, -1)
                cv2.circle(image, (76, 80), 8, 20, -1)
                cv2.circle(image, (124, 80), 8, 20, -1)
                cv2.line(image, (76, 125), (124, 125), 30, 5)
                ok, encoded = cv2.imencode(".png", image)
                self.assertTrue(ok)
                raw = encoded.tobytes()
                sample_path = Path(TEST_DIR.name) / f"sample-{index}.enc"
                sample_path.write_bytes(biometric_service.fernet.encrypt(raw))
                db.add(BiometricSample(id=str(uuid.uuid4()), person_id=person_id, encrypted_path=str(sample_path),
                                       sha256=hashlib.sha256(raw).hexdigest(), quality_score=95))
            db.commit()
            model = biometric_service.train(db)
            self.assertEqual(model["person_count"], 1)
            self.assertEqual(model["sample_count"], 15)
            model_path = Path(db.query(BiometricModel).filter_by(active=True).one().file_path)
            self.assertTrue(model_path.name.endswith(".yml.enc"))
            self.assertNotIn(b"%YAML", model_path.read_bytes())
            biometric_service._recognizer = None
            biometric_service._model_path = None
            recognizer, _ = biometric_service._active_recognizer(db)
            label, distance = recognizer.predict(image)
            self.assertEqual(label, person_id)
            self.assertLessEqual(distance, 65)

    def test_health_and_sf2_template(self) -> None:
        health = self.client.get("/api/health")
        self.assertEqual(health.status_code, 200, health.text)
        self.assertTrue(health.json()["ok"])
        template = self.client.get("/api/sf2/template", headers=self.headers)
        self.assertEqual(template.status_code, 200, template.text)
        self.assertTrue(template.json()["configured"])

    def test_extended_attendance_administration_and_disposal(self) -> None:
        from app.database import SessionLocal
        from app.models import AttendanceCorrection, AttendanceEvent, ExcusedAbsence, GradeScore, Person
        from app.services.attendance import MANILA, record_gate_match
        from app.services.settings_store import set_json

        created = self.client.post("/api/persons", headers=self.headers, json={
            "external_id": "EXTENDED-001", "lrn": "999999999981", "full_name": "REENTRY, LEARNER",
            "sex": "Male", "role": "Student", "grade": "10", "section": "Rizal",
            "guardian_phone": "+639181111111", "biometric_consent": False,
        })
        self.assertEqual(created.status_code, 200, created.text)
        person_id = created.json()["id"]
        with SessionLocal() as db:
            person = db.get(Person, person_id)
            set_json(db, "attendance", {"absence_cutoff": "09:00", "duplicate_cooldown_seconds": 0,
                                        "auto_close_enabled": True})
            directions = []
            for hour, minute in ((7, 20), (10, 0), (10, 20), (16, 15)):
                result = record_gate_match(db, person, 18.0, datetime(2026, 8, 18, hour, minute, tzinfo=MANILA))
                self.assertTrue(result["recorded"])
                directions.append(result["direction"])
            self.assertEqual(directions, ["Time In", "Time Out", "Time In", "Time Out"])

        holiday = self.client.post("/api/calendar/exceptions", headers=self.headers, json={
            "event_date": "2026-08-19", "event_type": "Holiday", "reason": "Authorized local holiday",
            "start_time": None, "end_time": None, "absence_cutoff": None,
        })
        self.assertEqual(holiday.status_code, 200, holiday.text)
        rows = self.client.get("/api/attendance?date=2026-08-19", headers=self.headers)
        self.assertEqual(next(item for item in rows.json() if item["person_id"] == person_id)["status"], "Holiday")

        excused = self.client.post("/api/attendance/excused", headers=self.headers, json={
            "person_id": person_id, "event_date": "2026-08-20", "reason": "Approved medical certificate MC-2026-08",
        })
        self.assertEqual(excused.status_code, 200, excused.text)
        rows = self.client.get("/api/attendance?date=2026-08-20", headers=self.headers)
        self.assertEqual(next(item for item in rows.json() if item["person_id"] == person_id)["status"], "Excused")

        correction = self.client.post("/api/attendance/corrections", headers=self.headers, json={
            "person_id": person_id, "attendance_date": "2026-08-18", "status": "Present",
            "time_in": "07:20", "time_out": "16:15", "reason": "Verified against the signed gate record",
        })
        self.assertEqual(correction.status_code, 200, correction.text)
        gradebook = self.client.put("/api/gradebook/disposal-gradebook", headers=self.headers, json={
            "class_key": "disposal-gradebook", "school_year": "2026-2027", "quarter": 1,
            "subject": "Mathematics", "change_reason": "Initial authorized score entry",
            "passing_grade": 75, "components": [{"id": "q1", "category": "Quizzes", "label": "Quiz 1", "weight": 100, "max_score": 10}],
            "scores": {str(person_id): {"q1": 9}},
        })
        self.assertEqual(gradebook.status_code, 200, gradebook.text)
        audit = self.client.get("/api/gradebook/disposal-gradebook/audit", headers=self.headers)
        self.assertEqual(audit.status_code, 200, audit.text)
        self.assertEqual(audit.json()[0]["reason"], "Initial authorized score entry")

        disposed = self.client.request("DELETE", f"/api/admin/persons/{person_id}/records", headers=self.headers, json={
            "reason": "Approved end-of-retention disposal request", "authorization_reference": "AUTH-DISPOSE-001",
            "confirmation": "EXTENDED-001",
        })
        self.assertEqual(disposed.status_code, 200, disposed.text)
        with SessionLocal() as db:
            self.assertIsNone(db.get(Person, person_id))
            self.assertEqual(db.query(AttendanceEvent).filter_by(person_id=person_id).count(), 0)
            self.assertEqual(db.query(AttendanceCorrection).filter_by(person_id=person_id).count(), 0)
            self.assertEqual(db.query(ExcusedAbsence).filter_by(person_id=person_id).count(), 0)
            self.assertEqual(db.query(GradeScore).filter_by(person_id=person_id).count(), 0)

        removed_exception = self.client.delete(f"/api/calendar/exceptions/{holiday.json()['id']}", headers=self.headers)
        self.assertEqual(removed_exception.status_code, 200, removed_exception.text)

    def test_multi_face_lbph_and_approved_roster_import(self) -> None:
        import cv2
        import numpy as np
        from openpyxl import Workbook, load_workbook
        from app.database import SessionLocal
        from app.models import BiometricSample, Person
        from app.services.biometrics import ProcessedFace, biometric_service
        from app.services.sf2 import generate_sf2

        people = []
        for external_id, name, sex in (("MULTI-001", "Zulu, Ana", "Female"), ("MULTI-002", "Alpha, Bea", "Female")):
            response = self.client.post("/api/persons", headers=self.headers, json={
                "external_id": external_id, "lrn": None, "full_name": name, "sex": sex, "role": "Student",
                "grade": "9", "section": "Multi", "guardian_phone": None, "biometric_consent": True,
            })
            self.assertEqual(response.status_code, 200, response.text)
            people.append(response.json())

        base_faces = []
        with SessionLocal() as db:
            for person_index, person in enumerate(people):
                for sample_index in range(15):
                    image = np.zeros((200, 200), dtype=np.uint8)
                    if person_index == 0:
                        for x in range(15 + sample_index % 3, 200, 30): cv2.line(image, (x, 0), (x, 199), 180, 12)
                    else:
                        for y in range(15 + sample_index % 3, 200, 30): cv2.line(image, (0, y), (199, y), 220, 12)
                    if sample_index == 0: base_faces.append(image.copy())
                    ok, encoded = cv2.imencode(".png", image); self.assertTrue(ok)
                    raw = encoded.tobytes(); path = Path(TEST_DIR.name) / f"multi-{person['id']}-{sample_index}.enc"
                    path.write_bytes(biometric_service.fernet.encrypt(raw))
                    db.add(BiometricSample(id=str(uuid.uuid4()), person_id=person["id"], encrypted_path=str(path),
                                           sha256=hashlib.sha256(raw).hexdigest(), quality_score=96))
            db.commit(); biometric_service.train(db)
            with patch.object(biometric_service, "process_faces", return_value=[
                ProcessedFace(base_faces[0], b"face-a", 96, (10, 10, 100, 100)),
                ProcessedFace(base_faces[1], b"face-b", 96, (120, 10, 100, 100)),
            ]):
                matches = biometric_service.recognize_many(db, b"two-face-frame")
            self.assertEqual({match[0].id for match in matches if match[0]}, {item["id"] for item in people})

        workbook = Workbook(); sheet = workbook.active
        sheet.append(["External ID", "Full Name", "Sex", "Role", "Grade", "Section", "LRN"])
        sheet.append(["ROSTER-002", "BETA, LEARNER", "Male", "Student", "8", "Roster", "888888888882"])
        sheet.append(["ROSTER-001", "ALPHA, LEARNER", "Male", "Student", "8", "Roster", "888888888881"])
        roster_stream = io.BytesIO(); workbook.save(roster_stream); roster_data = roster_stream.getvalue()
        files = {"file": ("approved-roster.xlsx", roster_data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        preview = self.client.post("/api/admin/roster/import", headers=self.headers, files=files,
                                   data={"approved_reference": "SJNHS-ROSTER-2026", "dry_run": "true"})
        self.assertEqual(preview.status_code, 200, preview.text); self.assertEqual(preview.json()["valid_rows"], 2)
        imported = self.client.post("/api/admin/roster/import", headers=self.headers, files=files,
                                    data={"approved_reference": "SJNHS-ROSTER-2026", "dry_run": "false"})
        self.assertEqual(imported.status_code, 200, imported.text); self.assertEqual(imported.json()["inserted"], 2)
        with SessionLocal() as db:
            sf2_path = generate_sf2(db, 2026, 8, "8", "Roster", "SJNHS", "2026-2027", "San Jose National High School")
        rendered = load_workbook(sf2_path, read_only=True, data_only=True)["School Form 2 (SF2)"]
        self.assertEqual(rendered["B14"].value, "ALPHA, LEARNER")
        self.assertEqual(rendered["B15"].value, "BETA, LEARNER")
        rendered.parent.close()

        for item in people:
            disposed = self.client.request("DELETE", f"/api/admin/persons/{item['id']}/records", headers=self.headers,
                                           json={"reason": "Test cleanup after multi-face verification",
                                                 "authorization_reference": "TEST-CLEANUP-001", "confirmation": item["external_id"]})
            self.assertEqual(disposed.status_code, 200, disposed.text)
        for external_id in ("ROSTER-001", "ROSTER-002"):
            item = next(row for row in self.client.get("/api/admin/persons", headers=self.headers).json() if row["external_id"] == external_id)
            self.client.request("DELETE", f"/api/admin/persons/{item['id']}/records", headers=self.headers,
                                json={"reason": "Test cleanup after roster verification",
                                      "authorization_reference": "TEST-CLEANUP-002", "confirmation": external_id})

    def test_z_secure_backup_and_restore_staging(self) -> None:
        passphrase = "Capstone-Recovery-2026!"
        created = self.client.post("/api/admin/backups", headers=self.headers, data={"passphrase": passphrase})
        self.assertEqual(created.status_code, 200, created.text)
        filename = created.json()["filename"]
        downloaded = self.client.get(f"/api/admin/backups/{filename}", headers=self.headers)
        self.assertEqual(downloaded.status_code, 200, downloaded.text)
        self.assertTrue(downloaded.content.startswith(b"EDUSCAN1"))
        restored = self.client.post("/api/admin/backups/restore", headers=self.headers,
                                    data={"passphrase": passphrase, "confirmation": "STAGE RESTORE"},
                                    files={"file": (filename, downloaded.content, "application/octet-stream")})
        self.assertEqual(restored.status_code, 200, restored.text)
        self.assertTrue(restored.json()["restart_required"])
        self.assertTrue((Path(TEST_DIR.name) / "pending_restore" / "ready.json").exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
