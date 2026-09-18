from __future__ import annotations

import hashlib
import html
import json
import re
import shutil
import uuid
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import EXPORT_DIR
from ..models import GeneratedReport, Person, User
from .attendance import active_people, raw_row
from .grading import gradebook_summary


REPORT_DIR = EXPORT_DIR / "registered"
REPORT_DIR.mkdir(parents=True, exist_ok=True)


def _safe_name(value: object) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "-", str(value)).strip("-._") or "report"


def register_report(db: Session, source: Path, report_type: str, parameters: dict, actor: User) -> GeneratedReport:
    if not source.exists():
        raise HTTPException(status_code=500, detail="Generated report file was not found")
    report_id = str(uuid.uuid4())
    target = REPORT_DIR / f"{report_id}-{source.name}"
    shutil.copy2(source, target)
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    record = GeneratedReport(
        id=report_id, report_type=report_type, filename=source.name, file_path=str(target), file_sha256=digest,
        parameters_json=json.dumps(parameters, default=str, sort_keys=True), status="Generated",
        generated_by=actor.id, generated_by_name=actor.full_name,
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


def _title(sheet, title: str, subtitle: str, last_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.cell(1, 1, title).font = Font(size=16, bold=True)
    sheet.cell(1, 1).alignment = Alignment(horizontal="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(2, 1, subtitle).alignment = Alignment(horizontal="center")


def _headers(sheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, 1):
        cell = sheet.cell(row, column, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def generate_grade_report_xlsx(db: Session, class_key: str) -> Path:
    summary = gradebook_summary(db, class_key)
    metadata, rows, statistics = summary["metadata"], summary["rows"], summary["statistics"]
    if not metadata["grade"] or not metadata["section"]:
        raise HTTPException(status_code=409, detail="Save the gradebook before generating its report")
    book = Workbook()
    sheet = book.active
    sheet.title = "Grade Summary"
    labels = ["LRN / ID", "Learner", "Initial Grade", "Transmuted Grade", "Result", "Completion"]
    _title(sheet, "EDUSCAN GRADING SUMMARY", f"SY {metadata['school_year']} · Quarter {metadata['quarter']} · {metadata['subject']} · Grade {metadata['grade']} {metadata['section']}", len(labels))
    _headers(sheet, 4, labels)
    for row_number, item in enumerate(rows, 5):
        values = [item["lrn"] or item["external_id"], item["full_name"], item["initial_grade"], item["transmuted_grade"], item["status"], "Complete" if item["complete"] else "Incomplete"]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
    stats_row = len(rows) + 7
    sheet.cell(stats_row, 1, "CLASS STATISTICS").font = Font(bold=True)
    stats = [("Learners", statistics["learners"]), ("Complete", statistics["complete"]), ("Incomplete", statistics["incomplete"]),
             ("Passed", statistics["passed"]), ("Below rule", statistics["below_rule"]), ("Average", statistics["average"]),
             ("Highest", statistics["highest"]), ("Lowest", statistics["lowest"])]
    for index, (label, value) in enumerate(stats, stats_row + 1):
        sheet.cell(index, 1, label).font = Font(bold=True)
        sheet.cell(index, 2, value)
    sheet.freeze_panes = "A5"
    for column, width in {"A": 20, "B": 34, "C": 16, "D": 18, "E": 16, "F": 14}.items():
        sheet.column_dimensions[column].width = width
    target = EXPORT_DIR / f"Grade-Summary-{_safe_name(metadata['school_year'])}-Q{metadata['quarter']}-{_safe_name(metadata['grade'])}-{_safe_name(metadata['section'])}-{_safe_name(metadata['subject'])}.xlsx"
    book.save(target)
    return target


def grade_report_html(db: Session, class_key: str) -> str:
    summary = gradebook_summary(db, class_key)
    metadata, rows, statistics = summary["metadata"], summary["rows"], summary["statistics"]
    body = "".join(
        f"<tr><td>{html.escape(item['lrn'] or item['external_id'])}</td><td>{html.escape(item['full_name'])}</td>"
        f"<td>{item['initial_grade'] if item['initial_grade'] is not None else '—'}</td><td>{item['transmuted_grade'] if item['transmuted_grade'] is not None else '—'}</td>"
        f"<td>{html.escape(item['status'])}</td></tr>" for item in rows
    )
    return _printable_html(
        "EduScan Grading Summary",
        f"SY {metadata['school_year']} · Quarter {metadata['quarter']} · {metadata['subject']} · Grade {metadata['grade']} {metadata['section']}",
        "<table><thead><tr><th>LRN / ID</th><th>Learner</th><th>Initial</th><th>Transmuted</th><th>Result</th></tr></thead><tbody>" + body + "</tbody></table>"
        f"<h2>Class statistics</h2><p>Learners: {statistics['learners']} · Complete: {statistics['complete']} · Incomplete: {statistics['incomplete']} · "
        f"Passed: {statistics['passed']} · Below rule: {statistics['below_rule']} · Average: {statistics['average'] if statistics['average'] is not None else '—'} · "
        f"Highest: {statistics['highest'] if statistics['highest'] is not None else '—'} · Lowest: {statistics['lowest'] if statistics['lowest'] is not None else '—'}</p>",
    )


def generate_gradebook_report_xlsx(db: Session, gradebook_id: int) -> Path:
    from .grading import get_gradebook_full
    full = get_gradebook_full(db, gradebook_id)
    if not full:
        raise HTTPException(status_code=404, detail="Gradebook not found")
        
    gb = full["gradebook"]
    rows = full["students"]
    statistics = full["statistics"]
    
    book = Workbook()
    sheet = book.active
    sheet.title = "Grade Summary"
    labels = ["LRN / ID", "Learner", "Initial Grade", "Transmuted Grade", "Result", "Completion"]
    _title(sheet, "EDUSCAN GRADING SUMMARY", f"SY {gb['school_year_name']} · Quarter {gb['quarter']} · {gb['subject_name']} · Grade {gb['grade_name']} {gb['section_name']}", len(labels))
    _headers(sheet, 4, labels)
    for row_number, item in enumerate(rows, 5):
        values = [item["lrn"] or item["external_id"], item["full_name"], item["initial_grade"], item["reported_grade"], item["status"], "Complete" if item["complete"] else "Incomplete"]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
            
    stats_row = len(rows) + 7
    sheet.cell(stats_row, 1, "CLASS STATISTICS").font = Font(bold=True)
    stats = [("Learners", statistics["learners"]), ("Complete", statistics["complete"]), ("Incomplete", statistics["incomplete"]),
             ("Passed", statistics["passed"]), ("Below rule", statistics["below_passing"]), ("Average", statistics["average"]),
             ("Highest", statistics["highest"]), ("Lowest", statistics["lowest"])]
    for index, (label, value) in enumerate(stats, stats_row + 1):
        sheet.cell(index, 1, label).font = Font(bold=True)
        sheet.cell(index, 2, value)
        
    sheet.freeze_panes = "A5"
    for column, width in {"A": 20, "B": 34, "C": 16, "D": 18, "E": 16, "F": 14}.items():
        sheet.column_dimensions[column].width = width
    target = EXPORT_DIR / f"Grade-Summary-{_safe_name(gb['school_year_name'])}-Q{gb['quarter']}-{_safe_name(gb['grade_name'])}-{_safe_name(gb['section_name'])}-{_safe_name(gb['subject_name'])}.xlsx"
    book.save(target)
    return target


def gradebook_report_html(db: Session, gradebook_id: int) -> str:
    from .grading import get_gradebook_full
    full = get_gradebook_full(db, gradebook_id)
    if not full:
        raise HTTPException(status_code=404, detail="Gradebook not found")
        
    gb = full["gradebook"]
    rows = full["students"]
    statistics = full["statistics"]
    
    body = "".join(
        f"<tr><td>{html.escape(item['lrn'] or item['external_id'])}</td><td>{html.escape(item['full_name'])}</td>"
        f"<td>{item['initial_grade'] if item['initial_grade'] is not None else '—'}</td><td>{item['reported_grade'] if item['reported_grade'] is not None else '—'}</td>"
        f"<td>{html.escape(item['status'])}</td></tr>" for item in rows
    )
    return _printable_html(
        "EduScan Grading Summary",
        f"SY {gb['school_year_name']} · Quarter {gb['quarter']} · {gb['subject_name']} · Grade {gb['grade_name']} {gb['section_name']}",
        "<table><thead><tr><th>LRN / ID</th><th>Learner</th><th>Initial</th><th>Transmuted</th><th>Result</th></tr></thead><tbody>" + body + "</tbody></table>"
        f"<h2>Class statistics</h2><p>Learners: {statistics['learners']} · Complete: {statistics['complete']} · Incomplete: {statistics['incomplete']} · "
        f"Passed: {statistics['passed']} · Below rule: {statistics['below_passing']} · Average: {statistics['average'] if statistics['average'] is not None else '—'} · "
        f"Highest: {statistics['highest'] if statistics['highest'] is not None else '—'} · Lowest: {statistics['lowest'] if statistics['lowest'] is not None else '—'}</p>",
    )



def attendance_range_rows(db: Session, starts_on: date, ends_on: date, role: str | None = None,
                          grade: str | None = None, section: str | None = None,
                          person_id: int | None = None) -> list[dict]:
    if ends_on < starts_on or (ends_on - starts_on).days > 366:
        raise HTTPException(status_code=422, detail="Attendance report range must be between 1 and 367 days")
    people = active_people(db, grade, section, students_only=bool(grade or section))
    if role:
        people = [item for item in people if item.role == role]
    if person_id:
        people = [item for item in people if item.id == person_id]
    result = []
    current = starts_on
    while current <= ends_on:
        for person in people:
            row = raw_row(db, person, current)
            row["attendance_date"] = current
            result.append(row)
        current += timedelta(days=1)
    return result


def generate_attendance_range_xlsx(db: Session, starts_on: date, ends_on: date, role: str | None = None,
                                   grade: str | None = None, section: str | None = None,
                                   person_id: int | None = None) -> Path:
    rows = attendance_range_rows(db, starts_on, ends_on, role, grade, section, person_id)
    book = Workbook()
    detail = book.active
    detail.title = "Attendance Detail"
    labels = ["Date", "ID / LRN", "Name", "Role", "Grade", "Section / assignment", "Time in", "Time out", "Status", "Source"]
    _title(detail, "EDUSCAN ATTENDANCE REPORT", f"{starts_on.isoformat()} through {ends_on.isoformat()}", len(labels))
    _headers(detail, 4, labels)
    for row_number, item in enumerate(rows, 5):
        values = [item["attendance_date"], item["lrn"] or item["external_id"], item["full_name"], item["role"], item["grade"] or "",
                  item["section"] or item["assignment"] or "", item["time_in"].strftime("%H:%M:%S") if item["time_in"] else "",
                  item["time_out"].strftime("%H:%M:%S") if item["time_out"] else "", item["status"], item["source"]]
        for column, value in enumerate(values, 1):
            detail.cell(row_number, column, value)
    detail.freeze_panes = "A5"
    for column, width in {"A": 13, "B": 20, "C": 34, "D": 24, "E": 10, "F": 24, "G": 12, "H": 12, "I": 16, "J": 30}.items():
        detail.column_dimensions[column].width = width

    summary_sheet = book.create_sheet("Summary")
    _title(summary_sheet, "ATTENDANCE SUMMARY", f"{starts_on.isoformat()} through {ends_on.isoformat()}", 9)
    summary_labels = ["ID / LRN", "Name", "Role", "Present", "Late", "Absent", "Excused", "No scan", "Time-out days"]
    _headers(summary_sheet, 4, summary_labels)
    grouped: dict[int, list[dict]] = defaultdict(list)
    for item in rows:
        grouped[item["person_id"]].append(item)
    for row_number, person_rows in enumerate(grouped.values(), 5):
        first = person_rows[0]
        counts = Counter(item["status"] for item in person_rows)
        values = [first["lrn"] or first["external_id"], first["full_name"], first["role"], counts["Present"], counts["Late"], counts["Absent"], counts["Excused"], counts["No scan"], counts["Time Out"]]
        for column, value in enumerate(values, 1):
            summary_sheet.cell(row_number, column, value)
    target = EXPORT_DIR / f"Attendance-{starts_on.isoformat()}-to-{ends_on.isoformat()}.xlsx"
    book.save(target)
    return target


def attendance_report_html(db: Session, starts_on: date, ends_on: date, role: str | None = None,
                           grade: str | None = None, section: str | None = None,
                           person_id: int | None = None) -> str:
    rows = attendance_range_rows(db, starts_on, ends_on, role, grade, section, person_id)
    grouped: dict[int, list[dict]] = defaultdict(list)
    for item in rows:
        grouped[item["person_id"]].append(item)
    body = ""
    for person_rows in grouped.values():
        first = person_rows[0]
        counts = Counter(item["status"] for item in person_rows)
        body += f"<tr><td>{html.escape(first['lrn'] or first['external_id'])}</td><td>{html.escape(first['full_name'])}</td><td>{html.escape(first['role'])}</td><td>{counts['Present']}</td><td>{counts['Late']}</td><td>{counts['Absent']}</td><td>{counts['Excused']}</td><td>{counts['No scan']}</td></tr>"
    return _printable_html("EduScan Attendance Summary", f"{starts_on.isoformat()} through {ends_on.isoformat()}",
                           "<table><thead><tr><th>ID / LRN</th><th>Name</th><th>Role</th><th>Present</th><th>Late</th><th>Absent</th><th>Excused</th><th>No scan</th></tr></thead><tbody>" + body + "</tbody></table>")


def _printable_html(title: str, subtitle: str, content: str) -> str:
    return f"""<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>
    body{{font-family:Arial,sans-serif;margin:32px;color:#111}}h1{{text-align:center;margin-bottom:4px}}.subtitle{{text-align:center;margin-bottom:24px}}
    table{{border-collapse:collapse;width:100%;font-size:12px}}th,td{{border:1px solid #777;padding:6px;text-align:left}}th{{background:#dbeafe}}
    @media print{{button{{display:none}}body{{margin:12mm}}}}button{{padding:8px 14px;margin-bottom:18px}}
    </style></head><body><button onclick='window.print()'>Print report</button><h1>{html.escape(title)}</h1><p class='subtitle'>{html.escape(subtitle)}</p>{content}</body></html>"""
