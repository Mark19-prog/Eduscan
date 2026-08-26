from __future__ import annotations

import hashlib
import io
import re
import uuid
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Person, RosterImportAudit, User


HEADER_ALIASES = {
    "external_id": {"external id", "school id", "student id", "employee id", "id"},
    "lrn": {"lrn", "learner reference number"},
    "full_name": {"full name", "name", "learner name", "student name"},
    "sex": {"sex", "gender"},
    "role": {"role", "personnel group", "type"},
    "grade": {"grade", "grade level"},
    "section": {"section"},
    "assignment": {"assignment", "office", "department"},
    "guardian_phone": {"guardian phone", "parent phone", "mobile number", "contact number"},
    "enrollment_status": {"enrollment status", "learner status", "movement status"},
    "enrollment_start_date": {"enrollment start", "enrollment start date", "transfer in date"},
    "enrollment_end_date": {"enrollment end", "enrollment end date", "transfer out date"},
    "transfer_school": {"transfer school", "previous school", "receiving school"},
}


def _normal(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _date_value(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ValueError("use YYYY-MM-DD") from exc


def parse_roster(data: bytes) -> tuple[list[dict], list[dict]]:
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Roster workbook exceeds 15 MB")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Roster workbook could not be read: {exc}") from exc
    mapped: dict[int, str] = {}
    for index, header in enumerate(headers):
        normalized = _normal(header)
        for field, aliases in HEADER_ALIASES.items():
            if normalized in aliases:
                mapped[index] = field
                break
    required = {"external_id", "full_name", "sex", "role"}
    if not required.issubset(set(mapped.values())):
        raise HTTPException(status_code=422, detail=f"Required roster columns: {', '.join(sorted(required))}")
    accepted: list[dict] = []
    errors: list[dict] = []
    for row_number, values in enumerate(rows, start=2):
        record = {}
        date_error = None
        for index, field in mapped.items():
            value = values[index] if index < len(values) else None
            if field in {"enrollment_start_date", "enrollment_end_date"}:
                try:
                    record[field] = _date_value(value)
                except ValueError as exc:
                    record[field] = None
                    date_error = f"{field.replace('_', ' ')}: {exc}"
            else:
                record[field] = str(value).strip() if value is not None else ""
        if not any(record.values()):
            continue
        record["sex"] = record.get("sex", "").title()
        role = record.get("role", "").casefold()
        record["role"] = {"student": "Student", "faculty": "Faculty", "teacher": "Faculty",
                          "non-teaching": "Non-teaching Personnel", "non teaching": "Non-teaching Personnel",
                          "non-teaching personnel": "Non-teaching Personnel"}.get(role, record.get("role", ""))
        status = record.get("enrollment_status", "").title() or "Regular"
        record["enrollment_status"] = status
        problems = [date_error] if date_error else []
        if len(record.get("external_id", "")) < 2: problems.append("missing school/employee ID")
        if len(record.get("full_name", "")) < 3: problems.append("missing full name")
        if record["sex"] not in {"Male", "Female"}: problems.append("sex must be Male or Female")
        if record["role"] not in {"Student", "Faculty", "Non-teaching Personnel"}: problems.append("invalid role")
        if record["role"] == "Student" and (not record.get("grade") or not record.get("section")):
            problems.append("student grade and section are required")
        if status not in {"Regular", "Transferred In", "Transferred Out"}:
            problems.append("enrollment status must be Regular, Transferred In, or Transferred Out")
        if record.get("enrollment_start_date") and record.get("enrollment_end_date") and record["enrollment_start_date"] > record["enrollment_end_date"]:
            problems.append("enrollment start date cannot be after enrollment end date")
        if problems:
            errors.append({"row": row_number, "errors": problems})
            continue
        record["lrn"] = record.get("lrn") or None
        record["grade"] = record.get("grade") or None
        record["section"] = record.get("section") or None
        record["assignment"] = record.get("assignment") or None
        record["guardian_phone"] = record.get("guardian_phone") or None
        record["transfer_school"] = record.get("transfer_school") or None
        accepted.append(record)
    workbook.close()
    return accepted, errors


def import_roster(db: Session, data: bytes, filename: str, approved_reference: str,
                  actor: User, dry_run: bool) -> dict:
    approved_reference = approved_reference.strip()
    if len(approved_reference) < 5:
        raise HTTPException(status_code=422, detail="Enter the school approval/document reference for this roster")
    rows, errors = parse_roster(data)
    inserted = updated = skipped = 0
    if dry_run:
        return {"dry_run": True, "valid_rows": len(rows), "invalid_rows": len(errors), "errors": errors[:100]}
    if errors:
        raise HTTPException(status_code=422, detail={"message": "Correct invalid roster rows before import", "errors": errors[:100]})
    for record in rows:
        person = db.scalar(select(Person).where(Person.external_id == record["external_id"]))
        if not person and record.get("lrn"):
            person = db.scalar(select(Person).where(Person.lrn == record["lrn"]))
        if person:
            for field, value in record.items():
                setattr(person, field, value)
            person.active = True
            updated += 1
        else:
            db.add(Person(**record, biometric_consent=False, active=True))
            inserted += 1
    db.add(RosterImportAudit(id=str(uuid.uuid4()), filename=filename[:255],
                             file_sha256=hashlib.sha256(data).hexdigest(), approved_reference=approved_reference,
                             inserted_count=inserted, updated_count=updated, skipped_count=skipped,
                             actor_user_id=actor.id, actor_name=actor.full_name))
    db.commit()
    return {"dry_run": False, "inserted": inserted, "updated": updated, "skipped": skipped, "invalid_rows": 0}
